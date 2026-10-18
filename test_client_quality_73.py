"""Current delivered proof must match the frozen approved presentation edition."""
import hashlib,json,subprocess
from pathlib import Path
from openpyxl import load_workbook
D=Path(__file__).parent/'docs'
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def test_current_proof_not_the_archived_baseline():
 data=json.loads((D/'assets/flagship/chart-data.json').read_text())
 assert data['source_product_sha']=='129f5f201ca13505f5ffc9b02ce40e7997011a61'
 assert data['presentation_edition']=='client-quality-2026-v3'
 assert data['workbook_sha256']==sha(D/'examples/harbourlink/cycle2/Premium_Workbook.xlsx')
 assert data['charts'][0]['values_by_series']==[[4,1,1,0]]
 assert data['charts'][2]['values_by_series']==[[1613,1,4,4,3]]
 assert data['charts'][3]['values_by_series']==[[264,0,1]]
 assert sum(sum(x) for x in data['charts'][1]['values_by_series'])==1622
 for cy in [1,2]:
  wb=load_workbook(D/f'examples/harbourlink/cycle{cy}/Premium_Workbook.xlsx')
  assert len(wb.sheetnames)==21 and len(wb['03 EXECUTIVE VIEW']._charts)==4
  html=(D/f'examples/harbourlink/cycle{cy}/Integrated_Management_Report.html').read_text()
  assert 'Confirm whether the missing governed relationship is expected or requires project-controls attention.' in html
  assert 'Use only conclusions that do not depend on these unresolved relationships.' in html

def test_film_preservation_and_retained_instrumental_audio():
 assert sha(D/'videos/wbslink-video-1-challenge.mp4')=='52ebc12339d3a26b20a734835cc3a52f2042de71159f8f586b643ce3ceee6538'
 for name in ['wbslink-video-2-how-it-works.mp4','wbslink-video-2-how-it-works-phone.mp4']:
  audio=subprocess.check_output(['ffmpeg','-v','error','-i',str(D/'videos'/name),'-map','0:a:0','-c','copy','-f','adts','-'])
  # Exact original AAC packet prefix; only the final packet beyond the new 80s cut is absent.
  retained=subprocess.check_output(['git','show','b9266f584ab9c8dd93753b3ba1c77c62936fc52d:docs/videos/wbslink-video-2-how-it-works.mp4'])
  import tempfile
  with tempfile.NamedTemporaryFile(suffix='.mp4') as f:
   f.write(retained);f.flush()
   old_audio=subprocess.check_output(['ffmpeg','-v','error','-i',f.name,'-map','0:a:0','-c','copy','-f','adts','-'])
  assert hashlib.sha256(old_audio).hexdigest()=='693bce75755a90ff8d976621d20f25f36ec40a1b714d5fe78ac3771ebc229acb'
  assert old_audio.startswith(audio) and len(old_audio)-len(audio)==445
 for vtt in ['wbslink-video-2-harbourlink.vtt','wbslink-video-2-how-it-works-phone.vtt']:
  text=(D/'videos'/vtt).read_text()
  assert text.count('-->')==10
  assert '1,622 current relationships plus three retired versions: 1,625 relationship versions in all.' in text
  assert 'Maya Chen' in text and 'does not individually confirm every relationship' in text
  assert 'Confirm whether the missing governed relationship is expected or requires project-controls attention.' in text

def test_recorded_and_live_journey_are_separate():
 html=(D/'index.html').read_text()
 assert 'Explore recorded June example' in html and 'Try the app' in html
 assert '1,614 remained valid' not in html
 guide=(D/'downloads/HarbourLink_First_Time_Guide.html').read_text()
 assert 'approve the structure' in guide.lower() and '1,356' in guide and '266' in guide
