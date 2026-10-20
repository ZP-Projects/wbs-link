"""Readable responsive web plots from the actual approved customer workbook cells.

Native chart images/workbook remain unchanged and directly accessible. This web
presentation has the same source-table counts, populations and series, not a new
model/score. Exact source labels are retained beside Canadian display aliases.
"""
from pathlib import Path
import argparse,hashlib,json,html
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

def cells(ws,ref):
 c0,r0,c1,r1=range_boundaries(ref.split('!')[1].replace('$',''));assert c0==c1
 return [ws.cell(r,c0).value for r in range(r0,r1+1)]

def build(workbook,terms,docs):
 cat=json.loads(terms.read_text());wb=load_workbook(workbook,data_only=False);ws=wb['03 EXECUTIVE VIEW'];assert len(wb.sheetnames)==21 and len(ws._charts)==4
 rows=[]
 for chart in ws._charts:
  labels=cells(ws,chart.series[0].cat.numRef.f);values=[cells(ws,s.val.numRef.f) for s in chart.series];assert all(isinstance(v,int) and v>=0 for x in values for v in x)
  rows.append({'source_categories':labels,'values_by_series':values,'display_categories':[cat['source_display_labels'].get(x,x) for x in labels],'category_ref':chart.series[0].cat.numRef.f,'series_refs':[s.val.numRef.f for s in chart.series]})
 assert rows[0]['values_by_series']==[[4,1,1,0]]
 assert rows[2]['values_by_series']==[[1613,1,4,4,3]]
 assert rows[3]['values_by_series']==[[264,0,1]]
 titles=['Owner-WBS coverage','Relationships by source','Reporting-cycle movement','Unresolved / outside WBS']
 scopes=['6 reported Owner-WBS branches. Coverage is not project health.','1,622 current relationships; retired excluded; split targets counted once.','1,625 relationships, retired included. Carry-forward and revalidation are distinct.','1,622 current relationships; retired excluded. UNSURE, NEW_SCOPE and NO_WBS remain distinct.']
 bits=[]
 for i,data in enumerate(rows):
  maximum=max(sum(vals) for vals in zip(*data['values_by_series'])) or 1;body=[]
  for j,(label,original) in enumerate(zip(data['display_categories'],data['source_categories'])):
   values=[x[j] for x in data['values_by_series']];bars=[]
   for k,value in enumerate(values):bars.append(f'<span class="plot-bar outcome-{k}" style="width:{value/maximum*100:.6f}%" aria-hidden="true"></span>')
   total=sum(values);original_note=f'<small class="original-chart-source">Source name as supplied: {html.escape(original)}</small>' if label!=original else ''
   numbers=' / '.join(f'{v:,}' for v in values)
   body.append(f'<div class="plot-row"><span class="plot-label">{html.escape(label)}{original_note}</span><span class="plot-track">{"".join(bars)}</span><strong class="plot-number">{numbers}</strong></div>')
  legend='<p class="plot-key"><span class="key-swatch outcome-0"></span>Related to WBS · <span class="key-swatch outcome-1"></span>Outside WBS · <span class="key-swatch outcome-2"></span>Unresolved. Values follow this order.</p>' if i==1 else ''
  bits.append(f'<figure class="web-chart" data-chart="{i+1}"><h4>{titles[i]}</h4><div class="live-plot" role="group" aria-label="{titles[i]}: actual approved workbook source-table data">{"".join(body)}</div>{legend}<figcaption><p class="chart-denominator">{scopes[i]}</p><a href="assets/flagship/chart{i+1}.png" target="_blank" rel="noopener" aria-label="Open chart {i+1} full size">Inspect original native workbook chart</a></figcaption></figure>')
 content='<div class="chart-proof-grid">'+''.join(bits)+'</div>'
 page=docs/'index.html';s=page.read_text();a=s.index('<div class="chart-proof-grid">');b=s.index('\n  <p class="evidence-note">',a);s=s[:a]+content+'</div>\n'+s[b:];page.write_text(s)
 payload={'copy_version':cat['version'],'source':'actual customer Premium Workbook /03 EXECUTIVE VIEW','governed_output_basis_sha':'72460cd9b78785607ccc255521c46dade0897bc6','workbook_sha256':hashlib.sha256(workbook.read_bytes()).hexdigest(),'charts':rows,'scope':'read-only responsive presentation; native chart objects, source rows and counting populations unchanged'}
 (docs/'assets/flagship/chart-data.json').write_text(json.dumps(payload,indent=2));wb.close();return payload

if __name__=='__main__':
 p=argparse.ArgumentParser();p.add_argument('--workbook',type=Path,required=True);p.add_argument('--terms',type=Path,required=True);p.add_argument('--docs',type=Path,required=True);a=p.parse_args();print(json.dumps(build(a.workbook,a.terms,a.docs),indent=2))
